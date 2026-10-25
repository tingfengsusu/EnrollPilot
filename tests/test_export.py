"""Excel 导出：结算单（含合计）、名单（集团 lookup 与筛选）、status_field 白名单。"""
from __future__ import annotations

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from openpyxl import load_workbook

from enroll.demo_store import DemoStore
from enroll.export import export_roster, export_settlement
from enroll.instructions import parse
from enroll.service import execute


def _tmp(name):
    return os.path.join(tempfile.mkdtemp(), name)


def _rows(path, sheet=None):
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    return list(ws.iter_rows(values_only=True))


def test_settlement_export():
    store = DemoStore().seed()
    path = export_settlement(store, _tmp("结算单.xlsx"))
    rows = _rows(path)
    assert rows[0][:12] == ("序号", "姓名", "公司", "人员类别", "工种", "培训类型",
                            "缴费", "公司收费", "返还", "缴费状态", "开票状态", "差异")
    assert rows[1][1] == "李蕾" and rows[1][6] == 700 and rows[1][7] == 400
    assert rows[1][11] in ("", None)                  # 拆分平 → 无差异（空串落盘读回 None）
    assert rows[-1][0] == "合计"                      # 两条种子：700/400/300 × 2
    assert rows[-1][6] == 1400 and rows[-1][7] == 800 and rows[-1][8] == 600


def test_settlement_flags_issue():
    store = DemoStore().seed()
    execute(store, parse("李蕾 已退"), operator="测试")
    path = export_settlement(store, _tmp("结算单2.xlsx"))
    rows = _rows(path)
    assert "已退但存在正向收款" in rows[1][11]


def test_roster_group_lookup_and_filter():
    store = DemoStore().seed()
    path = export_roster(store, _tmp("名单.xlsx"), group="某集团集团")
    rows = _rows(path)
    assert rows[0][5] == "集团"
    assert {r[1] for r in rows[1:]} == {"李蕾", "张莹"}   # 两家公司都属某集团集团
    assert all(r[5] == "某集团集团" for r in rows[1:])

    path2 = export_roster(store, _tmp("名单2.xlsx"),
                          status_field="缴费状态", status_value="已缴费")
    rows2 = _rows(path2)
    assert [r[1] for r in rows2[1:]] == ["李蕾"]

    path3 = export_roster(store, _tmp("名单3.xlsx"), period="6月21日")
    assert len(_rows(path3)) == 3                     # 表头 + 2 条


def test_status_field_whitelist():
    store = DemoStore().seed()
    with pytest.raises(ValueError):
        export_roster(store, _tmp("x.xlsx"), status_field="姓名", status_value="李蕾")


if __name__ == "__main__":
    test_settlement_export()
    test_settlement_flags_issue()
    test_roster_group_lookup_and_filter()
    test_status_field_whitelist()
    print("ok: test_export")
